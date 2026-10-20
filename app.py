import base64
import io
import json

import pandas as pd
import streamlit as st
from openai import OpenAI

st.set_page_config(page_title="Mon Claude", page_icon="💬")

MODEL = st.secrets["MODEL"]
SYSTEM = (
    "Tu es un assistant personnel. Réponds toujours en français, de façon claire et soignée. "
    "Quand on te demande un fichier Excel ou Word, utilise l'outil create_excel ou create_word "
    "(ne colle jamais le contenu complet du fichier dans ta réponse, résume-le simplement). "
    "Pour un Excel : mets des en-têtes clairs, des nombres en vrais nombres, et des formules "
    "(texte commençant par =) quand c'est utile."
)

# ---------------- Mot de passe ----------------
if not st.session_state.get("ok"):
    pwd = st.text_input("Mot de passe", type="password")
    if pwd:
        if pwd == st.secrets["APP_PASSWORD"]:
            st.session_state.ok = True
            st.rerun()
        else:
            st.error("Mot de passe incorrect")
    st.stop()

client = OpenAI(api_key=st.secrets["API_KEY"], base_url=st.secrets["BASE_URL"])
ss = st.session_state
ss.setdefault("messages", [{"role": "system", "content": SYSTEM}])
ss.setdefault("display", [])

# ---------------- Outils de création de fichiers ----------------
TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "create_excel",
            "description": "Crée un fichier Excel (.xlsx) avec une ou plusieurs feuilles.",
            "parameters": {
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "Nom du fichier, ex: pronostics.xlsx"},
                    "sheets": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string"},
                                "headers": {"type": "array", "items": {"type": "string"}},
                                "rows": {"type": "array", "items": {"type": "array", "items": {}}},
                            },
                            "required": ["name", "headers", "rows"],
                        },
                    },
                },
                "required": ["filename", "sheets"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_word",
            "description": "Crée un document Word (.docx). Le contenu est en Markdown simple (# titres, ## sous-titres, - listes, paragraphes).",
            "parameters": {
                "type": "object",
                "properties": {
                    "filename": {"type": "string"},
                    "content": {"type": "string"},
                },
                "required": ["filename", "content"],
            },
        },
    },
]


def build_excel(args):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for sh in args["sheets"]:
        ws = wb.create_sheet(str(sh["name"])[:31] or "Feuille")
        ws.append(sh["headers"])
        for row in sh["rows"]:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_word(args):
    from docx import Document

    doc = Document()
    for line in args["content"].splitlines():
        s = line.rstrip()
        if s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.lstrip().startswith(("- ", "* ")):
            doc.add_paragraph(s.lstrip()[2:], style="List Bullet")
        elif s:
            doc.add_paragraph(s)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def run_tool(name, args, files):
    try:
        if name == "create_excel":
            fname = args["filename"] if args["filename"].endswith(".xlsx") else args["filename"] + ".xlsx"
            files.append({"name": fname, "data": build_excel(args)})
            return f"Fichier {fname} créé."
        if name == "create_word":
            fname = args["filename"] if args["filename"].endswith(".docx") else args["filename"] + ".docx"
            files.append({"name": fname, "data": build_word(args)})
            return f"Fichier {fname} créé."
        return "Outil inconnu."
    except Exception as e:
        return f"Erreur : {e}"


# ---------------- Lecture des fichiers envoyés ----------------
def read_upload(up):
    name = up.name.lower()
    data = up.getvalue()
    try:
        if name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
            b64 = base64.b64encode(data).decode()
            return {"type": "image_url", "image_url": {"url": f"data:{up.type};base64,{b64}"}}
        if name.endswith((".xlsx", ".xlsm", ".xls")):
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
            text = "\n\n".join(f"[Feuille {n}]\n{df.to_csv(index=False)}" for n, df in sheets.items())
        elif name.endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
        elif name.endswith(".docx"):
            from docx import Document
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as e:
        text = f"(lecture impossible : {e})"
    return {"type": "text", "text": f"--- Fichier joint : {up.name} ---\n{text[:80000]}"}


# ---------------- Appel au modèle ----------------
def ask(content):
    ss.messages.append({"role": "user", "content": content})
    files, used_model, final = [], MODEL, ""
    for _ in range(6):
        resp = client.chat.completions.create(
            model=MODEL, messages=ss.messages, tools=TOOLS, max_tokens=8000
        )
        used_model = resp.model or MODEL
        msg = resp.choices[0].message
        entry = {"role": "assistant", "content": msg.content or ""}
        if msg.tool_calls:
            entry["tool_calls"] = [
                {"id": t.id, "type": "function",
                 "function": {"name": t.function.name, "arguments": t.function.arguments}}
                for t in msg.tool_calls
            ]
        ss.messages.append(entry)
        if not msg.tool_calls:
            final = msg.content or ""
            break
        for t in msg.tool_calls:
            try:
                args = json.loads(t.function.arguments or "{}")
            except Exception:
                args = {}
            ss.messages.append(
                {"role": "tool", "tool_call_id": t.id, "content": run_tool(t.function.name, args, files)}
            )
    return final, files, used_model


# ---------------- Interface ----------------
with st.sidebar:
    if st.button("🗑️ Nouvelle conversation"):
        ss.messages = [{"role": "system", "content": SYSTEM}]
        ss.display = []
        st.rerun()
    st.caption(f"Modèle demandé : {MODEL}")

st.title("💬 Mon Claude")

for i, m in enumerate(ss.display):
    with st.chat_message(m["role"]):
        st.markdown(m["text"])
        if m.get("model"):
            st.caption(f"Modèle : {m['model']}")
        for j, f in enumerate(m.get("files", [])):
            st.download_button(f"⬇️ {f['name']}", f["data"], file_name=f["name"], key=f"d{i}_{j}")

chat = st.chat_input("Écris ton message… (+ pour joindre des fichiers)", accept_file="multiple")
if chat:
    prompt = chat.text or "Voici mes fichiers."
    content = [{"type": "text", "text": prompt}] + [read_upload(u) for u in (chat.files or [])]
    note = "".join(f"\n📎 {u.name}" for u in (chat.files or []))
    ss.display.append({"role": "user", "text": prompt + note})
    with st.chat_message("user"):
        st.markdown(prompt + note)
    with st.chat_message("assistant"):
        with st.spinner("Claude réfléchit…"):
            try:
                text, files, used_model = ask(content)
            except Exception as e:
                text, files, used_model = f"⚠️ Erreur : {e}", [], None
        st.markdown(text)
    ss.display.append({"role": "assistant", "text": text, "files": files, "model": used_model})
    st.rerun()
